import csv
import io

import pytest
from fastapi import HTTPException
from invoicelens.db import SessionLocal
from invoicelens.models import Document, Job
from invoicelens.service import get_export, refresh_findings
from invoicelens.workflow import run_job
from openpyxl import load_workbook
from reportlab.pdfgen import canvas


def invoice_pdf(
    number: str = "INV-901", po_number: str = "PO-901", total: str = "111.00"
) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(612, 792))
    y = 740
    for line in [
        "INVOICE",
        "Vendor: Test Fasteners",
        f"Invoice #: {number} Date: 06/16/2026",
        f"PO #: {po_number} Currency: USD Due Date: 07/16/2026",
        "SKU DESCRIPTION QTY UNIT UNIT PRICE AMOUNT",
        "SKU-1 Stainless fastener 2 ea 50.00 100.00",
        "Subtotal: 100.00",
        "Tax: 10.00",
        "Freight: 0.00",
        "Discount: 0.00",
        f"Total: {total}",
    ]:
        pdf.drawString(50, y, line)
        y -= 20
    pdf.save()
    return buffer.getvalue()


def po_pdf() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(612, 792))
    y = 740
    for line in [
        "PURCHASE ORDER",
        "Vendor: Test Fasteners",
        "PO #: PO-901 Currency: USD Delivery Date: 07/16/2026",
        "SKU DESCRIPTION QTY UNIT UNIT PRICE AMOUNT",
        "SKU-1 Stainless fastener 2 ea 50.00 100.00",
        "Subtotal: 100.00",
        "Tax: 0.00",
        "Freight: 0.00",
        "Discount: 0.00",
        "Total: 100.00",
    ]:
        pdf.drawString(50, y, line)
        y -= 20
    pdf.save()
    return buffer.getvalue()


def test_title_does_not_replace_explicit_invoice_number():
    from invoicelens.extraction import DemoExtractionAdapter, PageContent

    page = PageContent(
        1,
        "INVOICE LARCHSTONE COMPONENTS\nInvoice #: 017-26-0001 Date: 03.06.2026",
        "native",
        612,
        792,
        [],
    )
    extracted = DemoExtractionAdapter().extract([page], b"", "application/pdf")
    assert extracted.number == "017-26-0001"


def test_auth_upload_review_export_and_workspace_scope(client, operator_headers):
    invalid = client.post(
        "/api/auth/login", json={"email": "operator@example.com", "password": "wrong"}
    )
    assert invalid.status_code == 401
    assert (
        client.get("/api/auth/me", headers={"Authorization": "Bearer !!!!"}).status_code
        == 401
    )

    content = invoice_pdf()
    upload = client.post(
        "/api/uploads",
        headers=operator_headers,
        files=[("files", ("../../unsafe.pdf", content, "application/pdf"))],
    )
    assert upload.status_code == 200, upload.text
    item = upload.json()["items"][0]
    document_id, job_id = item["document_id"], item["job_id"]
    assert item["status"] == "queued"
    duplicate = client.post(
        "/api/uploads",
        headers=operator_headers,
        files=[("files", ("different.pdf", content, "application/pdf"))],
    )
    assert duplicate.json()["items"][0]["document_id"] == document_id
    assert duplicate.json()["items"][0]["deduplicated"] is True
    run_job(job_id)

    detail = client.get(f"/api/documents/{document_id}", headers=operator_headers)
    assert detail.status_code == 200, detail.text
    document = detail.json()
    assert document["filename"] == "unsafe.pdf"
    assert document["fields"]["vendor"]["value"] == "Test Fasteners"
    assert document["fields"]["vendor"]["bbox"] is not None
    assert document["lines"][0]["evidence"]["unit_price"]["bbox"] is not None
    assert any(f["code"] == "TOTAL_MISMATCH" for f in document["findings"])
    source_response = client.get(
        f"/api/documents/{document_id}/source", headers=operator_headers
    )
    assert source_response.content == content
    assert source_response.headers["Cache-Control"] == "no-store"

    other = client.post(
        "/api/auth/login",
        json={"email": "operator2@example.com", "password": "DemoPass123!"},
    )
    assert other.status_code == 200
    other_headers = {}
    assert (
        client.get(f"/api/documents/{document_id}", headers=other_headers).status_code
        == 404
    )
    assert (
        client.get(
            f"/api/documents/{document_id}/source", headers=other_headers
        ).status_code
        == 404
    )
    assert client.get(f"/api/jobs/{job_id}", headers=other_headers).status_code == 404

    restored = client.post(
        "/api/auth/login",
        json={"email": "operator@example.com", "password": "DemoPass123!"},
    )
    assert restored.status_code == 200

    early = client.post(
        f"/api/documents/{document_id}/approve",
        headers=operator_headers,
        json={"version": document["version"]},
    )
    assert early.status_code == 409
    warning_id = next(
        f["id"] for f in document["findings"] if f["code"] == "PO_NOT_FOUND"
    )
    decision_path = f"/api/documents/{document_id}/findings/{warning_id}/decisions"
    accepted_warning = client.post(
        decision_path, headers=operator_headers, json={"decision": "accepted"}
    )
    assert accepted_warning.status_code == 200, accepted_warning.text
    rejected_warning = client.post(
        decision_path, headers=operator_headers, json={"decision": "rejected"}
    )
    assert rejected_warning.status_code == 200, rejected_warning.text
    assert any(
        item["before"] == {"status": "accepted"}
        and item["after"] == {"status": "rejected"}
        for item in rejected_warning.json()["history"]
    )
    same_decision = client.post(
        decision_path, headers=operator_headers, json={"decision": "rejected"}
    )
    assert same_decision.status_code == 200, same_decision.text
    assert len(same_decision.json()["history"]) == len(
        rejected_warning.json()["history"]
    )
    corrected = client.post(
        f"/api/documents/{document_id}/corrections",
        headers=operator_headers,
        json={"field": "total", "value": "110.00", "reason": "OCR digit corrected"},
    )
    assert corrected.status_code == 200, corrected.text
    document = corrected.json()
    assert document["version"] == 2
    assert document["fields"]["total"]["raw"] == "111.00"
    assert document["fields"]["total"]["value"] == "110.00"
    assert document["fields"]["total"]["bbox"] is not None
    assert not any(f["code"] == "TOTAL_MISMATCH" for f in document["findings"])
    corrected = client.post(
        f"/api/documents/{document_id}/corrections",
        headers=operator_headers,
        json={
            "field": "vendor",
            "value": "=SUM(1,1)",
            "reason": "Formula escape check",
        },
    )
    assert corrected.status_code == 200, corrected.text
    version = corrected.json()["version"]
    approved = client.post(
        f"/api/documents/{document_id}/approve",
        headers=operator_headers,
        json={"version": version},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "approved"
    csv_response = client.get(
        f"/api/documents/{document_id}/exports?format=csv", headers=operator_headers
    )
    assert csv_response.status_code == 200, csv_response.text
    assert csv_response.headers["Cache-Control"] == "no-store"
    rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert rows[1][rows[0].index("vendor")] == "'=SUM(1,1)"
    assert rows[1][rows[0].index("total")] == "110.00"
    assert (
        client.get(
            f"/api/documents/{document_id}/exports?format=csv", headers=operator_headers
        ).content
        == csv_response.content
    )
    workbook = client.get(
        f"/api/documents/{document_id}/exports?format=xlsx", headers=operator_headers
    )
    assert workbook.status_code == 200
    assert workbook.content.startswith(b"PK")
    sheet = load_workbook(io.BytesIO(workbook.content), read_only=True).active
    headings = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    assert values[headings.index("schema_version")] == "1.0"
    assert values[headings.index("total")] == "110.00"
    assert values[headings.index("line_unit_price")] == "50.00"
    vendor_cell = sheet.cell(2, headings.index("vendor") + 1)
    assert vendor_cell.value == "'=SUM(1,1)"
    assert vendor_cell.data_type == "s"
    structured = client.get(
        f"/api/documents/{document_id}/exports?format=json", headers=operator_headers
    )
    assert structured.json()["schema_version"] == "1.0"

    changed = client.post(
        f"/api/documents/{document_id}/corrections",
        headers=operator_headers,
        json={
            "field": "lines.0.unit_price",
            "value": "49.00",
            "reason": "new evidence",
        },
    )
    assert changed.status_code == 200
    assert changed.json()["approval"] is None
    assert changed.json()["lines"][0]["evidence"]["unit_price"]["raw"] == "50.00"
    assert (
        client.get(
            f"/api/documents/{document_id}/exports?format=csv", headers=operator_headers
        ).status_code
        == 409
    )
    restored_line = client.post(
        f"/api/documents/{document_id}/corrections",
        headers=operator_headers,
        json={
            "field": "lines.0.unit_price",
            "value": "50.00",
            "reason": "Recheck against source",
        },
    )
    assert restored_line.status_code == 200, restored_line.text
    latest_version = restored_line.json()["version"]
    assert latest_version == 5
    reapproved = client.post(
        f"/api/documents/{document_id}/approve",
        headers=operator_headers,
        json={"version": latest_version},
    )
    assert reapproved.status_code == 200, reapproved.text
    assert reapproved.json()["approval"]["version"] == latest_version
    assert (
        client.get(
            f"/api/documents/{document_id}/exports?format=csv", headers=operator_headers
        ).status_code
        == 200
    )

    # A PO arriving later changes findings for the same extraction version.
    po_upload = client.post(
        "/api/uploads",
        headers=operator_headers,
        files=[("files", ("late-po.pdf", po_pdf(), "application/pdf"))],
    )
    assert po_upload.status_code == 200, po_upload.text
    run_job(po_upload.json()["items"][0]["job_id"])
    after_po = client.get(f"/api/documents/{document_id}", headers=operator_headers)
    assert after_po.status_code == 200, after_po.text
    assert after_po.json()["version"] == latest_version
    assert after_po.json()["approval"] is None
    assert (
        after_po.json()["matches"][0]["document_id"]
        == po_upload.json()["items"][0]["document_id"]
    )
    assert any(
        f["code"] == "PO_SELECTION_REQUIRED" for f in after_po.json()["findings"]
    )
    selected = client.post(
        f"/api/documents/{document_id}/matches/decision",
        headers=operator_headers,
        json={
            "version": latest_version,
            "decision": "selected",
            "po_id": after_po.json()["matches"][0]["po_id"],
            "note": "Reviewed vendor mismatch against source",
        },
    )
    assert selected.status_code == 200, selected.text
    selected_detail = selected.json()
    assert selected_detail["version"] == latest_version + 1
    assert selected_detail["match_resolution"]["decision"] == "selected"
    assert selected_detail["matches"][0]["status"] == "selected"
    vendor_mismatch_id = next(
        f["id"]
        for f in selected_detail["findings"]
        if f["code"] == "PO_VENDOR_MISMATCH"
    )
    accepted_mismatch = client.post(
        f"/api/documents/{document_id}/findings/{vendor_mismatch_id}/decisions",
        headers=operator_headers,
        json={"decision": "accepted", "note": "Approved as reviewed test record"},
    )
    assert accepted_mismatch.status_code == 200, accepted_mismatch.text
    approved_again = client.post(
        f"/api/documents/{document_id}/approve",
        headers=operator_headers,
        json={"version": selected_detail["version"]},
    )
    assert approved_again.status_code == 200, approved_again.text
    assert approved_again.json()["approval"]["version"] == selected_detail["version"]


def test_upload_validation(client, operator_headers):
    unsupported = client.post(
        "/api/uploads",
        headers=operator_headers,
        files=[("files", ("bad.txt", b"not a document", "text/plain"))],
    )
    assert unsupported.json()["items"][0]["status"] == "rejected"
    damaged = client.post(
        "/api/uploads",
        headers=operator_headers,
        files=[("files", ("bad.pdf", b"%PDF-1.7 broken", "application/pdf"))],
    )
    assert "damaged" in damaged.json()["items"][0]["error"].lower()


def test_review_decision_invalidates_approval_and_incomplete_jobs_cannot_approve(
    client, operator_headers
):
    content = invoice_pdf("INV-REVIEW-901", "PO-REVIEW-901", "110.00")
    upload = client.post(
        "/api/uploads",
        files=[("files", ("finding-review.pdf", content, "application/pdf"))],
    )
    assert upload.status_code == 200, upload.text
    item = upload.json()["items"][0]
    document_id, job_id = item["document_id"], item["job_id"]
    run_job(job_id)
    detail_path = f"/api/documents/{document_id}"
    document = client.get(detail_path).json()
    assert document["version"] == 1
    finding_id = next(
        f["id"] for f in document["findings"] if f["code"] == "PO_NOT_FOUND"
    )
    approval_path = f"{detail_path}/approve"
    export_path = f"{detail_path}/exports?format=json"
    decision_path = f"{detail_path}/findings/{finding_id}/decisions"

    with SessionLocal() as db:
        job = db.get(Job, job_id)
        assert job is not None
        job.status = "processing"
        db.commit()
    assert client.post(approval_path, json={"version": 1}).status_code == 409
    with SessionLocal() as db:
        job = db.get(Job, job_id)
        document_row = db.get(Document, document_id)
        assert job is not None and document_row is not None
        job.status = "complete"
        document_row.status = "validating"
        db.commit()
    assert client.post(approval_path, json={"version": 1}).status_code == 409
    with SessionLocal() as db:
        document_row = db.get(Document, document_id)
        assert document_row is not None
        refresh_findings(db, document_row)
    approved = client.post(approval_path, json={"version": 1})
    assert approved.status_code == 200, approved.text
    assert client.get(export_path).status_code == 200

    client.post(
        "/api/auth/login",
        json={"email": "viewer@example.com", "password": "DemoPass123!"},
    )
    assert client.post(decision_path, json={"decision": "accepted"}).status_code == 403
    client.post(
        "/api/auth/login",
        json={"email": "operator2@example.com", "password": "DemoPass123!"},
    )
    assert client.post(decision_path, json={"decision": "accepted"}).status_code == 404
    client.post(
        "/api/auth/login",
        json={"email": "operator@example.com", "password": "DemoPass123!"},
    )

    with SessionLocal() as stale_db:
        stale_document = stale_db.get(Document, document_id)
        assert stale_document is not None and stale_document.approved_version == 1
        accepted = client.post(decision_path, json={"decision": "accepted"})
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()["version"] == 1
        assert accepted.json()["approval"] is None
        assert accepted.json()["status"] == "needs_review"
        assert client.get(export_path).status_code == 409
        with pytest.raises(HTTPException) as stale_export:
            get_export(stale_db, stale_document, "json")
        assert stale_export.value.status_code == 409

    assert client.post(approval_path, json={"version": 1}).status_code == 200
    rejected = client.post(decision_path, json={"decision": "rejected"})
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["approval"] is None
    assert any(
        event["before"] == {"status": "accepted"}
        and event["after"] == {"status": "rejected"}
        for event in rejected.json()["history"]
    )
    assert client.get(export_path).status_code == 409
    assert client.post(approval_path, json={"version": 1}).status_code == 200
    repeated = client.post(decision_path, json={"decision": "rejected"})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["approval"]["version"] == 1
    assert client.get(export_path).status_code == 200

    corrected = client.post(
        f"{detail_path}/corrections",
        json={"field": "date", "value": "06/17/2026", "reason": "Source date reviewed"},
    )
    assert corrected.status_code == 200, corrected.text
    assert corrected.json()["version"] == 2
    with SessionLocal() as db:
        document_row = db.get(Document, document_id)
        assert document_row is not None
        document_row.status = "validating"
        db.commit()
    assert client.post(approval_path, json={"version": 2}).status_code == 409
    with SessionLocal() as db:
        document_row = db.get(Document, document_id)
        assert document_row is not None
        refresh_findings(db, document_row)
    reapproved = client.post(approval_path, json={"version": 2})
    assert reapproved.status_code == 200, reapproved.text
    assert reapproved.json()["approval"]["version"] == 2
